"""Word nativo con identidad y logos integrados, sin solicitudes externas."""
from io import BytesIO
import math
import re
from urllib.parse import urlparse
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
from django.utils import timezone
from .presentacion import datos_manual, logo_disponible
from .composicion_documento import componer_manual
from .tablas_documento import anchos_normas

ANCHO_DOCUMENTO_MM = 186


def _imagen(parrafo, identificador, ancho, alto):
    logo = logo_disponible(identificador) if identificador else None
    if not logo:
        return
    with logo.imagen.open('rb') as archivo:
        data = archivo.read()
    with Image.open(BytesIO(data)) as imagen:
        ratio = min(ancho / imagen.width, alto / imagen.height)
        parrafo.add_run().add_picture(BytesIO(data), width=Mm(imagen.width * ratio), height=Mm(imagen.height * ratio))


def _borde_tabla(tabla, color):
    bordes = OxmlElement('w:tblBorders')
    linea = OxmlElement('w:bottom')
    for clave, valor in [('val', 'single'), ('sz', '12'), ('color', color.lstrip('#'))]:
        linea.set(qn('w:' + clave), valor)
    bordes.append(linea)
    tabla._tbl.tblPr.append(bordes)


def _cabecera(contenedor, datos, compacto=False):
    bloques = datos['bloques']
    kwargs = {'rows': 1, 'cols': len(bloques)}
    if compacto:
        kwargs['width'] = Mm(ANCHO_DOCUMENTO_MM)
    tabla = contenedor.add_table(**kwargs)
    tabla.autofit = False
    for celda, bloque in zip(tabla.rows[0].cells, bloques):
        celda.width = Mm(ANCHO_DOCUMENTO_MM / len(bloques))
        p = celda.paragraphs[0]
        if bloque.get('logo'):
            _imagen(p, bloque['logo'], 26, 9)
            if p._p.xpath('.//w:drawing'):
                p.add_run('  ')
        if len(bloques) > 1:
            r = p.add_run(bloque['rol'].upper() + ' · ')
            r.font.size = Pt(7); r.font.color.rgb = RGBColor.from_string('64748B')
        r = p.add_run(bloque['nombre']); r.bold = True; r.font.size = Pt(11)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        if bloque['lineas']:
            p = celda.add_paragraph(' · '.join(f"{l['etiqueta']}: {l['valor']}" for l in bloque['lineas']))
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            for r in p.runs:
                r.font.size = Pt(8)
    _borde_tabla(tabla, datos['color'])


def _margen_superior(datos):
    # Reserva ajustada al texto real; un encabezado sin logo deja de ocupar 4 cm.
    ancho = ANCHO_DOCUMENTO_MM / len(datos['bloques'])
    alturas = []
    for bloque in datos['bloques']:
        logo = logo_disponible(bloque.get('logo')) if bloque.get('logo') else None
        disponibles = max(24, ancho - (28 if logo else 0))
        nombres = max(1, math.ceil(len(bloque['nombre']) * 1.8 / disponibles))
        detalles = ' · '.join(f"{l['etiqueta']}: {l['valor']}" for l in bloque['lineas'])
        lineas = math.ceil(len(detalles) * 1.35 / ancho) if detalles else 0
        altura = max(9 if logo else 0, nombres * 4.2 + (3 if len(datos['bloques']) > 1 else 0))
        alturas.append(altura + lineas * 3.2 + 2)
    return max(20, math.ceil(8 + max(alturas) + 4))


def _inline(parrafo, nodo, negrita=False, cursiva=False, codigo=False, sub=False, sup=False):
    if isinstance(nodo, NavigableString):
        run = parrafo.add_run(str(nodo) if codigo else re.sub(r'\s+', ' ', str(nodo)))
        run.bold, run.italic = negrita, cursiva
        run.font.subscript, run.font.superscript = sub, sup
        if codigo:
            run.font.name = 'Consolas'
        return
    if not isinstance(nodo, Tag):
        return
    if nodo.name == 'br':
        parrafo.add_run().add_break(); return
    if nodo.name == 'img':
        if nodo.get('alt'):
            parrafo.add_run(nodo['alt'])
        return
    if nodo.name == 'a' and urlparse(nodo.get('href', '')).scheme in ('https', 'http', 'mailto'):
        enlace = OxmlElement('w:hyperlink')
        enlace.set(qn('r:id'), parrafo.part.relate_to(nodo['href'], RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r'); texto = OxmlElement('w:t'); texto.text = nodo.get_text(); run.append(texto); enlace.append(run)
        parrafo._p.append(enlace)
        return
    for hijo in nodo.children:
        _inline(parrafo, hijo, negrita or nodo.name in ('b', 'strong'),
                cursiva or nodo.name in ('em', 'i'), codigo or nodo.name == 'code',
                sub or nodo.name == 'sub', sup or nodo.name == 'sup')


def _numero_lista(doc, ordenada):
    numbering = doc.part.numbering_part.element
    ids = [int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))]
    abstract_id = max(ids, default=-1) + 1
    abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), str(abstract_id))
    level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
    for tag, value in [('start', '1'), ('numFmt', 'decimal' if ordenada else 'bullet'), ('lvlText', '%1.' if ordenada else '•')]:
        elem = OxmlElement('w:' + tag); elem.set(qn('w:val'), value); level.append(elem)
    abstract.append(level); numbering.append(abstract)
    return numbering.add_num(abstract_id).numId


def _lista(doc, nodo, nivel=0):
    num_id = _numero_lista(doc, nodo.name == 'ol')
    for item in nodo.find_all('li', recursive=False):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(6 + nivel * 6)
        p.paragraph_format.first_line_indent = Mm(-4)
        num = p._p.get_or_add_pPr().get_or_add_numPr()
        num.get_or_add_ilvl().val = 0; num.get_or_add_numId().val = num_id
        for hijo in item.children:
            if isinstance(hijo, Tag) and hijo.name in ('ol', 'ul'):
                _lista(doc, hijo, nivel + 1)
            else:
                _inline(p, hijo)


def _span(celda, atributo):
    try:
        return max(1, min(30, int(celda.get(atributo, 1))))
    except (ValueError, TypeError):
        return 1


def _tabla(doc, nodo):
    filas = [r for r in nodo.find_all('tr') if r.find_parent('table') == nodo]
    if not filas:
        return
    posiciones, usados, columnas = [], set(), 0
    for y, fila in enumerate(filas):
        x = 0
        for celda in fila.find_all(['td', 'th'], recursive=False):
            while (y, x) in usados:
                x += 1
            ancho, alto = _span(celda, 'colspan'), min(_span(celda, 'rowspan'), len(filas) - y)
            posiciones.append((y, x, alto, ancho, celda))
            usados.update((yy, xx) for yy in range(y, y + alto) for xx in range(x, x + ancho))
            x += ancho; columnas = max(columnas, x)
    if not columnas:
        return
    tabla = doc.add_table(rows=len(filas), cols=columnas)
    tabla.style = 'Table Grid'
    encabezados = [c.get_text(' ', strip=True) for c in filas[0].find_all(['td', 'th'], recursive=False)]
    anchos = anchos_normas(encabezados)
    if anchos and columnas == len(anchos):
        tabla.autofit = False
        for columna, porcentaje in zip(tabla.columns, anchos):
            columna.width = Mm(ANCHO_DOCUMENTO_MM * porcentaje / 100)
        for fila in tabla.rows:
            for celda, porcentaje in zip(fila.cells, anchos):
                celda.width = Mm(ANCHO_DOCUMENTO_MM * porcentaje / 100)
    for y, x, alto, ancho, contenido in posiciones:
        celda = tabla.cell(y, x)
        if ancho > 1 or alto > 1:
            celda = celda.merge(tabla.cell(y + alto - 1, x + ancho - 1))
        celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _inline(celda.paragraphs[0], contenido, negrita=contenido.name == 'th')
        if anchos:
            margenes = OxmlElement('w:tcMar')
            for lado in ('left', 'right'):
                margen = OxmlElement('w:' + lado); margen.set(qn('w:w'), '90'); margen.set(qn('w:type'), 'dxa'); margenes.append(margen)
            celda._tc.get_or_add_tcPr().append(margenes)
            for parrafo in celda.paragraphs:
                parrafo.paragraph_format.space_after = Pt(3)
                parrafo.paragraph_format.line_spacing = 1.1
            if x == 0 and (contenido.name == 'th' or re.fullmatch(r'(?:NIIF|NIF|NIC|IFRS|IAS)\s*[A-Z]?[ .-]?\d+(?:[ .-]\d+)*', contenido.get_text(' ', strip=True), re.I)):
                celda._tc.get_or_add_tcPr().append(OxmlElement('w:noWrap'))
    if filas[0].find('th'):
        repeat = OxmlElement('w:tblHeader'); tabla.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph()


def _contenido(doc, nodo):
    for hijo in nodo.children:
        if isinstance(hijo, NavigableString):
            if hijo.strip():
                doc.add_paragraph(str(hijo))
            continue
        nombre = hijo.name
        if nombre in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            p = doc.add_heading(level=int(nombre[1])); _inline(p, hijo)
        elif nombre in ('ol', 'ul'):
            _lista(doc, hijo)
        elif nombre == 'table':
            _tabla(doc, hijo)
        elif nombre == 'pre':
            p = doc.add_paragraph(); r = p.add_run(hijo.get_text()); r.font.name = 'Consolas'; r.font.size = Pt(9)
        elif nombre in ('div', 'section'):
            _contenido(doc, hijo)
        else:
            p = doc.add_paragraph(style='Quote' if nombre == 'blockquote' else None); _inline(p, hijo)


def generar_docx(manual):
    datos = datos_manual(manual)
    composicion = componer_manual(manual)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(12)
    section.top_margin, section.bottom_margin = Mm(_margen_superior(datos)), Mm(18)
    section.header_distance, section.footer_distance = Mm(8), Mm(10)
    normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    doc.styles['Title'].font.size = Pt(18)
    doc.styles['Title'].paragraph_format.line_spacing = 1.05
    for nombre, tamano in [('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 11)]:
        estilo = doc.styles[nombre]
        estilo.font.size = Pt(tamano)
        estilo.paragraph_format.space_before = Pt(8)
        estilo.paragraph_format.space_after = Pt(4)
    rgb = [int(datos['color'][i:i + 2], 16) / 255 for i in (1, 3, 5)]
    luminancia = sum(p * (v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4) for p, v in zip([0.2126, 0.7152, 0.0722], rgb))
    color_titulo = datos['color'].lstrip('#') if 1.05 / (luminancia + 0.05) >= 4.5 else '172b45'
    for heading in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[heading].font.color.rgb = RGBColor.from_string(color_titulo)
    section.header.paragraphs[0].paragraph_format.space_after = Pt(0)
    section.header.paragraphs[0].paragraph_format.line_spacing = Pt(1)
    _cabecera(section.header, datos, compacto=True)
    empresa_visible = any(b['nombre'] == datos['empresa_nombre'] for b in datos['bloques'])
    if not empresa_visible:
        doc.add_paragraph('Empresa: ' + datos['empresa_nombre'])
    titulo = doc.add_heading(composicion['titulo'], 0)
    titulo.paragraph_format.space_before = Pt(0)
    titulo.paragraph_format.space_after = Pt(5)
    fecha = timezone.localtime(manual.fecha_creacion).strftime('%d/%m/%Y')
    metadata = doc.add_paragraph(f"{fecha} · {'Aprobado' if manual.aprobado else 'Pendiente de aprobación'}")
    metadata.paragraph_format.space_after = Pt(6)
    for run in metadata.runs:
        run.font.size = Pt(8)
    if composicion['alcance']:
        alcance = doc.add_paragraph(' · '.join(f"{c['clave']} — {c['descripcion']}" for c in composicion['alcance']))
        alcance.paragraph_format.space_after = Pt(6)
        for run in alcance.runs:
            run.font.size = Pt(9)
    _contenido(doc, BeautifulSoup(composicion['contenido_html'], 'html.parser'))
    pie = section.footer.paragraphs[0]; pie.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pie.add_run('Página ')
    for instruccion, separador in [('PAGE', ' de '), ('NUMPAGES', '')]:
        campo = OxmlElement('w:fldSimple'); campo.set(qn('w:instr'), instruccion); pie._p.append(campo)
        if separador:
            pie.add_run(separador)
    actualizar = OxmlElement('w:updateFields'); actualizar.set(qn('w:val'), 'true'); doc.settings.element.append(actualizar)
    doc.core_properties.author = next((b['nombre'] for b in datos['bloques'] if b['rol'] == 'Elaborado por'), '')
    doc.core_properties.title = manual.titulo
    output = BytesIO(); doc.save(output)
    return output.getvalue()
