from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from unittest.mock import patch
import importlib

from PIL import Image
from docx import Document
from django.apps import apps
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, override_settings
from django.urls import reverse
from cuentas.models import Tenant, MembresiaTenant, IdentidadDespacho, PresentacionEmpresa, LogoDocumento
from cuentas.marca_forms import normalizar_logo
from .models import ManualContable, CambioPresentacionManual, UsoIA
from .presentacion import construir_presentacion, datos_manual


def imagen(nombre='logo.png', size=(280, 90), color=(20, 90, 120, 160)):
    output = BytesIO()
    Image.new('RGBA', size, color).save(output, 'PNG')
    return SimpleUploadedFile(nombre, output.getvalue(), content_type='image/png')


@override_settings(DEBUG=True, SECURE_SSL_REDIRECT=False)
class PresentacionTests(TestCase):
    def setUp(self):
        self.directorio = TemporaryDirectory()
        self.addCleanup(self.directorio.cleanup)
        self.config = override_settings(PRIVATE_MEDIA_ROOT=Path(self.directorio.name))
        self.config.enable(); self.addCleanup(self.config.disable)
        self.owner = User.objects.create_user(username='titular', email='personal-titular@example.com')
        self.otro = User.objects.create_user(username='otro-despacho')
        self.contador = User.objects.create_user(username='contador', email='personal-lector@example.com')
        self.lector = User.objects.create_user(username='lector')
        self.empresa = Tenant.objects.create(nombre='Cliente Uno', propietario=self.owner, rfc='XAXX010101000', ciudad='Mérida')
        self.empresa2 = Tenant.objects.create(nombre='Cliente Dos', propietario=self.otro)
        for user, tenant, rol in [(self.owner, self.empresa, 'propietario'), (self.otro, self.empresa2, 'propietario'),
                                  (self.contador, self.empresa, 'contador'), (self.contador, self.empresa2, 'contador'),
                                  (self.lector, self.empresa, 'solo_lectura')]:
            MembresiaTenant.objects.create(usuario=user, tenant=tenant, rol=rol)
        self.logo1 = self.logo(self.owner)
        self.logo2 = self.logo(self.owner, self.empresa, size=(80, 240), color=(40, 170, 100, 255))
        self.despacho = IdentidadDespacho.objects.create(titular=self.owner, nombre='Despacho Alfa', logo=self.logo1,
                         correo='contacto@alfa.example', color='#18547a')
        IdentidadDespacho.objects.create(titular=self.otro, nombre='Despacho Beta')
        self.config_empresa = PresentacionEmpresa.objects.create(organizacion=self.empresa, logo=self.logo2)
        self.manual = ManualContable.objects.create(tenant=self.empresa, titulo='Políticas de tesorería', creado_por=self.owner,
            contenido='# Tesorería\n\n**Control de caja** y conciliaciones.\n\n1. Preparar\n2. Revisar\n\n- Archivo\n    - Documentos\n\n| Cuenta | Control |\n|---|---|\n| Bancos | Conciliar |\n\n<script>alert(1)</script>')
        self.login(self.owner, self.empresa)

    def login(self, user, empresa=None):
        self.client.force_login(user)
        session = self.client.session
        if empresa:
            session['tenant_id'] = empresa.pk
        else:
            session.pop('tenant_id', None)
        session.save()

    def logo(self, owner, empresa=None, **kwargs):
        logo = LogoDocumento(propietario=owner, organizacion=empresa, ambito='empresa' if empresa else 'despacho')
        data = normalizar_logo(imagen(**kwargs))
        logo.imagen.save(data.name, data, save=True)
        return logo

    def test_identidad_depende_del_propietario_y_no_del_lector(self):
        self.login(self.contador, self.empresa)
        response = self.client.get(reverse('detalle_manual', args=[self.manual.pk]))
        self.assertContains(response, 'Despacho Alfa')
        self.assertNotContains(response, 'Despacho Beta')
        documento = Document(BytesIO(self.client.get(reverse('exportar_manual', args=[self.manual.pk])).content))
        self.assertEqual(documento.core_properties.author, 'Despacho Alfa')
        xml = documento.part.element.xml
        self.assertNotIn(self.contador.email, xml)
        self.assertNotIn(self.owner.email, xml)

    def test_composiciones_y_datos_visibles(self):
        for modo, esperado in [('ambos', ['Despacho Alfa', 'Cliente Uno']), ('empresa', ['Cliente Uno']), ('despacho', ['Despacho Alfa'])]:
            self.config_empresa.composicion = modo; self.config_empresa.save()
            datos = construir_presentacion(Tenant.objects.get(pk=self.empresa.pk))
            self.assertEqual([b['nombre'] for b in datos['bloques']], esperado)
            self.assertEqual(datos['empresa_nombre'], 'Cliente Uno')
        self.config_empresa.campos_empresa = []; self.config_empresa.composicion = 'empresa'; self.config_empresa.save()
        self.assertEqual(construir_presentacion(Tenant.objects.get(pk=self.empresa.pk))['bloques'][0]['lineas'], [])
        self.config_empresa.logo = None; self.config_empresa.save()
        self.assertIsNone(construir_presentacion(Tenant.objects.get(pk=self.empresa.pk))['bloques'][0]['logo'])

    def test_borradores_cambian_aprobados_conservan_y_reaprobar_no_altera(self):
        url = reverse('aprobar_manual', args=[self.manual.pk])
        self.client.post(url)
        self.manual.refresh_from_db()
        original = self.manual.presentacion_snapshot
        self.despacho.nombre = 'Despacho Nuevo'; self.despacho.logo = self.logo(self.owner); self.despacho.save()
        self.empresa.nombre = 'Cliente Renombrado'; self.empresa.save()
        self.client.post(url)
        self.manual.refresh_from_db()
        self.assertEqual(self.manual.presentacion_snapshot, original)
        self.assertEqual(datos_manual(self.manual)['bloques'][0]['logo'], str(self.logo1.pk))
        self.assertTrue(Path(self.logo1.imagen.path).exists())
        self.assertEqual(self.client.get(reverse('logo_documento', args=[self.logo1.pk])).status_code, 200)
        borrador = ManualContable.objects.create(tenant=Tenant.objects.get(pk=self.empresa.pk), titulo='Otro', contenido='Otro', creado_por=self.owner)
        self.assertEqual(datos_manual(borrador)['bloques'][0]['nombre'], 'Despacho Nuevo')

    def test_actualizacion_explicita_previsualiza_audita_y_no_consume_ia(self):
        self.client.post(reverse('aprobar_manual', args=[self.manual.pk]))
        self.despacho.nombre = 'Despacho Renovado'; self.despacho.save()
        url = reverse('actualizar_presentacion_manual', args=[self.manual.pk])
        response = self.client.get(url)
        self.assertContains(response, 'Despacho Renovado')
        with patch('catalogo.ia.OpenAI') as api:
            response = self.client.post(url, {'version_previa': response.context['version_previa']})
            api.assert_not_called()
        self.assertEqual(response.status_code, 302)
        self.manual.refresh_from_db()
        self.assertTrue(self.manual.aprobado)
        self.assertEqual(self.manual.contenido.splitlines()[0], '# Tesorería')
        self.assertEqual(self.manual.presentacion_snapshot['bloques'][0]['nombre'], 'Despacho Renovado')
        self.assertEqual(CambioPresentacionManual.objects.get().autor, self.owner)
        self.assertFalse(UsoIA.objects.exists())

    def test_vista_previa_caducada_o_manipulada_no_actualiza(self):
        self.client.post(reverse('aprobar_manual', args=[self.manual.pk]))
        url = reverse('actualizar_presentacion_manual', args=[self.manual.pk])
        token = self.client.get(url).context['version_previa']
        self.despacho.nombre = 'Cambió mientras revisabas'; self.despacho.save()
        self.assertEqual(self.client.post(url, {'version_previa': token}).status_code, 409)
        self.assertEqual(self.client.post(url, {'version_previa': 'falso'}).status_code, 409)
        self.assertFalse(CambioPresentacionManual.objects.exists())

    def test_permisos_de_empresa_y_marcas_ajenas(self):
        self.login(self.contador, self.empresa)
        self.assertEqual(self.client.get(reverse('presentacion_empresa', args=[self.empresa.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse('logo_documento', args=[self.logo1.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse('logo_documento', args=[self.logo2.pk])).status_code, 200)
        self.login(self.lector, self.empresa)
        self.assertEqual(self.client.get(reverse('presentacion_empresa', args=[self.empresa.pk])).status_code, 403)
        self.assertEqual(self.client.post(reverse('aprobar_manual', args=[self.manual.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse('actualizar_presentacion_manual', args=[self.manual.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse('presentacion_empresa', args=[self.empresa2.pk])).status_code, 403)
        otra_empresa = Tenant.objects.create(nombre='Otro cliente Alfa', propietario=self.owner)
        logo_ajeno = self.logo(self.owner, otra_empresa)
        self.assertEqual(self.client.get(reverse('logo_documento', args=[logo_ajeno.pk])).status_code, 404)
        otra_empresa.delete()
        self.assertEqual(self.client.get(reverse('logo_documento', args=[logo_ajeno.pk])).status_code, 404)
        MembresiaTenant.objects.filter(usuario=self.lector).delete()
        self.assertEqual(self.client.get(reverse('logo_documento', args=[self.logo1.pk])).status_code, 404)

    def test_mi_despacho_solo_modifica_identidad_propia_y_funciona_sin_empresa_activa(self):
        self.login(self.contador)
        url = reverse('mi_despacho')
        self.assertEqual(self.client.get(url).status_code, 200)
        self.client.post(url, {'nombre': 'Marca del contador', 'color': '#123456', 'titular': self.owner.pk, 'accion': 'guardar'})
        self.despacho.refresh_from_db()
        self.assertEqual(self.despacho.nombre, 'Despacho Alfa')
        self.assertEqual(IdentidadDespacho.objects.get(titular=self.contador).nombre, 'Marca del contador')

    def test_preview_logo_se_conserva_al_guardar_sin_volver_a_subirlo(self):
        url = reverse('mi_despacho')
        response = self.client.post(url, {'nombre': 'Despacho con logo', 'color': '#123456',
                   'archivo_logo': imagen(), 'accion': 'preview'})
        self.assertEqual(response.status_code, 200)
        token = response.context['form']['logo_previo'].value()
        self.assertTrue(token)
        self.despacho.refresh_from_db()
        self.assertEqual(self.despacho.nombre, 'Despacho Alfa')
        self.assertEqual(self.despacho.logo, self.logo1)
        response = self.client.post(url, {'nombre': 'Despacho con logo', 'color': '#123456',
                                   'logo_previo': token, 'accion': 'guardar'})
        self.assertEqual(response.status_code, 302)
        self.despacho.refresh_from_db()
        self.assertNotEqual(self.despacho.logo_id, self.logo1.pk)
        self.assertTrue(Path(self.logo1.imagen.path).exists())
        self.assertTrue(Path(self.despacho.logo.imagen.path).exists())
        with self.assertRaises(ValueError):
            _ = self.despacho.logo.imagen.url

    def test_preview_falso_no_publica_logo(self):
        self.client.post(reverse('mi_despacho'), {'nombre': 'Intento', 'color': '#123456',
                        'logo_previo': 'falso', 'accion': 'guardar'})
        self.despacho.refresh_from_db()
        self.assertEqual(self.despacho.nombre, 'Despacho Alfa')

    def test_validacion_de_imagenes_y_color(self):
        url = reverse('mi_despacho')
        for archivo in [SimpleUploadedFile('logo.svg', b'<svg/>', content_type='image/png'),
                        SimpleUploadedFile('logo.png', b'0' * (2 * 1024 * 1024 + 1)),
                        imagen(size=(4097, 2))]:
            response = self.client.post(url, {'nombre': 'Alfa', 'color': '#123456', 'archivo_logo': archivo, 'accion': 'guardar'})
            self.assertTrue(response.context['form'].errors.get('archivo_logo'))
        response = self.client.post(url, {'nombre': 'Alfa', 'color': 'red;position:fixed', 'accion': 'guardar'})
        self.assertTrue(response.context['form'].errors.get('color'))
        clean = normalizar_logo(imagen(size=(80, 240)))
        with Image.open(clean) as image:
            self.assertEqual(image.size, (80, 240))
            self.assertEqual(image.mode, 'RGBA')
            self.assertFalse(image.info)

    def test_docx_autocontenido_listas_tablas_y_logos_repetidos(self):
        response = self.client.get(reverse('exportar_manual', args=[self.manual.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response['Content-Disposition'].endswith('.docx"'))
        doc = Document(BytesIO(response.content))
        self.assertGreaterEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[-1].cell(1, 0).text, 'Bancos')
        with ZipFile(BytesIO(response.content)) as zip:
            media = [f for f in zip.namelist() if f.startswith('word/media/')]
            self.assertEqual(len(media), 2)
            header = zip.read('word/header1.xml').decode()
            self.assertIn('Despacho Alfa', header)
            self.assertIn('Cliente Uno', header)
            xml = zip.read('word/document.xml').decode()
            self.assertIn('Control de caja', xml)
            self.assertIn('w:numPr', xml)
            self.assertIn('w:b', xml)
            self.assertNotIn('<script>', xml)
            self.assertNotIn('Tokens registrados', xml)
            rels = zip.read('word/_rels/header1.xml.rels').decode()
            self.assertNotIn('TargetMode="External"', rels)

    def test_migracion_fija_aprobados_anteriores_sin_logos(self):
        ManualContable.objects.filter(pk=self.manual.pk).update(aprobado=True, presentacion_snapshot={})
        modulo = importlib.import_module('catalogo.migrations.0013_manualcontable_presentacion_snapshot_and_more')
        modulo.conservar_aprobados(apps, type('Editor', (), {'connection': connection})())
        self.manual.refresh_from_db()
        self.assertEqual(self.manual.presentacion_snapshot['version'], 0)
        self.assertIsNone(self.manual.presentacion_snapshot['bloques'][0]['logo'])
        self.empresa.nombre = 'Cambió el nombre'; self.empresa.save()
        self.assertEqual(datos_manual(self.manual)['empresa_nombre'], 'Cliente Uno')

    def test_aprobacion_y_actualizacion_funcionan_si_el_autor_original_fue_eliminado(self):
        ManualContable.objects.filter(pk=self.manual.pk).update(creado_por=None)
        self.assertEqual(self.client.post(reverse('aprobar_manual', args=[self.manual.pk])).status_code, 302)
        self.despacho.nombre = 'Despacho con nueva identidad'; self.despacho.save()
        url = reverse('actualizar_presentacion_manual', args=[self.manual.pk])
        token = self.client.get(url).context['version_previa']
        self.assertEqual(self.client.post(url, {'version_previa': token}).status_code, 302)
        self.manual.refresh_from_db()
        self.assertIsNone(self.manual.creado_por)
        self.assertTrue(self.manual.aprobado)
        self.assertEqual(self.manual.presentacion_snapshot['bloques'][0]['nombre'], 'Despacho con nueva identidad')

    def test_salida_compacta_sin_duplicar_identidad_ni_titulo_y_sin_modificar_texto(self):
        from .models import CodigoAgrupador
        from .composicion_documento import componer_manual
        self.config_empresa.logo = None
        self.config_empresa.composicion = 'empresa'
        self.config_empresa.save()
        codigo = CodigoAgrupador.objects.create(nivel=3, codigoagrupador='602.41', nombre='Honorarios a personas morales')
        self.manual.codigos_sat.add(codigo)
        contenido = ('## Manual Contable — Cuenta 602.41\n\n'
                     '"Honorarios a personas morales"\n\n'
                     '**Código SAT:** 602.41 — Honorarios a personas morales\n\n'
                     '**Entidad:** Cliente Uno **Giro:** Comercial\n\n'
                     '**Nota técnica preliminar:** conservar íntegra esta explicación.\n\n'
                     '## Referencias\n\nReferencia normativa por verificar.')
        self.manual.titulo = 'Manual: 602.41'
        self.manual.contenido = contenido
        self.manual.save()
        salida = componer_manual(self.manual)
        self.assertEqual(salida['titulo'], 'Manual Contable — Cuenta 602.41')
        self.assertNotIn('Manual Contable — Cuenta', salida['contenido_html'])
        self.assertEqual(salida['contenido_html'].count('Honorarios a personas morales'), 1)
        self.assertIn('conservar íntegra esta explicación', salida['contenido_html'])
        self.assertEqual(salida['alcance'], [])
        doc = Document(BytesIO(self.client.get(reverse('exportar_manual', args=[self.manual.pk])).content))
        texto = '\n'.join(p.text for p in doc.paragraphs)
        self.assertEqual(texto.count('Manual Contable — Cuenta 602.41'), 1)
        self.assertNotIn('Empresa: Cliente Uno', texto)
        self.assertNotIn('Preparado para', texto)
        self.assertLessEqual(doc.sections[0].top_margin.mm, 25)
        self.manual.refresh_from_db()
        self.assertEqual(self.manual.contenido, contenido)

    def test_titulos_personalizados_y_secciones_tecnicas_se_conservan(self):
        from .composicion_documento import componer_manual
        self.manual.titulo = 'Políticas autorizadas del despacho'
        self.manual.contenido = '# Tesorería\n\n## Riesgos específicos\n\nNo eliminar ninguna observación.'
        salida = componer_manual(self.manual)
        self.assertEqual(salida['titulo'], self.manual.titulo)
        self.assertIn('Tesorería', salida['contenido_html'])
        self.assertIn('Riesgos específicos', salida['contenido_html'])
        self.assertIn('No eliminar ninguna observación.', salida['contenido_html'])

    def test_tabla_de_normas_reserva_ancho_para_codigo_y_titulo(self):
        from .composicion_documento import componer_manual
        from bs4 import BeautifulSoup
        self.manual.contenido = ('## NIIF/IFRS aplicables\n\n'
            '| NIIF | TÍTULO | APLICACIÓN A LA CUENTA 602.41 |\n|---|---|---|\n'
            '| NIIF 15 | Ingresos de contratos con clientes | Explicación extensa. ' + 'Texto de apoyo. ' * 30 + '|\n'
            '| NIC 24 | Revelación de partes relacionadas | Revisar los documentos. |')
        self.manual.save()
        salida = BeautifulSoup(componer_manual(self.manual)['contenido_html'], 'html.parser')
        tabla = salida.select_one('.document-standards-table')
        self.assertIsNotNone(tabla)
        self.assertEqual([c['style'] for c in tabla.select('colgroup col')], ['width: 13%', 'width: 29%', 'width: 58%'])
        self.assertIn('Revelación de partes relacionadas', tabla.get_text())
        self.assertEqual(tabla.parent.get('role'), 'region')
        doc = Document(BytesIO(self.client.get(reverse('exportar_manual', args=[self.manual.pk])).content))
        tabla_word = doc.tables[0]
        anchos = [c.width.mm for c in tabla_word.columns]
        self.assertGreater(anchos[0], 23)
        self.assertGreater(anchos[1], 52)
        self.assertGreater(anchos[2], anchos[1])
        self.assertAlmostEqual(doc.sections[0].left_margin.mm, 12, places=1)
        self.assertIn('w:noWrap', tabla_word.cell(1, 0)._tc.xml)
        self.assertEqual(tabla_word.cell(1, 0).text, 'NIIF 15')
        self.manual.refresh_from_db()
        self.assertEqual(self.manual.contenido.count('Texto de apoyo.'), 30)

    def test_tablas_de_cuentas_t_no_se_confunden_con_tablas_de_normas(self):
        from .tablas_documento import anchos_normas
        self.assertIsNone(anchos_normas(['Debe', 'Haber']))
        self.assertIsNone(anchos_normas(['Cuenta', 'Debe', 'Haber']))
        self.assertEqual(anchos_normas(['NIIF/IFRS', 'Título', 'Aplicación']), [13, 29, 58])
