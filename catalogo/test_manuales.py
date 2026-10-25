from unittest.mock import patch

from django.test import TestCase, override_settings, Client
from django.urls import reverse

from cuentas.models import Tenant
from .models import CatalogoCuentas, ManualContable
from . import tests as catalogo_tests
from .views import buscar_manual_existente


@override_settings(DEBUG=True, SECURE_SSL_REDIRECT=False)
class GestionManualesTests(TestCase):
    setUp = catalogo_tests.CatalogoCuentaDeleteTests.setUp

    def manual(self, cuentas=None, tenant=None, **kwargs):
        manual = ManualContable.objects.create(
            tenant=tenant or self.tenant, titulo=kwargs.pop('titulo', 'Manual de caja'),
            contenido=kwargs.pop('contenido', '# Procedimiento\n\nControl del efectivo.'),
            creado_por=self.user, **kwargs)
        manual.cuentas.set([self.cuenta] if cuentas is None else cuentas)
        return manual

    def generar(self, cuentas):
        return self.client.post(reverse('generar_manual'), {'cuentas': cuentas})

    def test_no_genera_de_nuevo_el_mismo_alcance(self):
        manual = self.manual()
        with patch('catalogo.views.generar_contenido_manual') as generar:
            response = self.generar([self.cuenta.pk])
        self.assertRedirects(response, reverse('detalle_manual', args=[manual.pk]))
        generar.assert_not_called()
        self.assertEqual(self.tenant.manuales.count(), 1)

    def test_rechaza_varias_cuentas_aunque_exista_un_manual_de_ese_alcance(self):
        otra = CatalogoCuentas.objects.create(tenant=self.tenant, cuenta='1001', descripcion='Otra caja', codigoagrupador=self.codigo)
        manual = self.manual([self.cuenta, otra])
        with patch('catalogo.views.generar_contenido_manual') as generar:
            response = self.generar([otra.pk, self.cuenta.pk, otra.pk])
        self.assertRedirects(response, reverse('generar_manual'))
        generar.assert_not_called()
        self.assertEqual(self.tenant.manuales.count(), 1)

    def test_rechaza_dos_cuentas_validas_sin_consumir_ia(self):
        otra = CatalogoCuentas.objects.create(tenant=self.tenant, cuenta='1001', descripcion='Otra caja', codigoagrupador=self.codigo)
        with patch('catalogo.views.generar_contenido_manual') as generar, patch('panel.controls.reservar_generacion') as reservar:
            response = self.generar([self.cuenta.pk, otra.pk])
        self.assertRedirects(response, reverse('generar_manual'))
        generar.assert_not_called(); reservar.assert_not_called()
        self.assertFalse(self.tenant.manuales.exists())

    def test_rechaza_dos_cuentas_sat_sin_consumir_ia(self):
        from .models import CodigoAgrupador
        otra = CodigoAgrupador.objects.create(nivel=1, codigoagrupador='102', nombre='Bancos')
        self.cuenta.delete()
        with patch('catalogo.views.generar_contenido_manual') as generar:
            response = self.generar([f'sat:{self.codigo.pk}', f'sat:{otra.pk}'])
        self.assertRedirects(response, reverse('generar_manual'))
        generar.assert_not_called()
        self.assertFalse(self.tenant.manuales.exists())

    def test_funcion_de_generacion_rechaza_alcance_multiple(self):
        from .views import generar_contenido_manual, ManualGenerationError
        with patch('catalogo.views.generar_texto') as ia:
            with self.assertRaises(ManualGenerationError):
                generar_contenido_manual([self.cuenta, self.cuenta], self.tenant)
        ia.assert_not_called()

    def test_cuenta_cubierta_por_manual_conjunto_no_se_duplica(self):
        otra = CatalogoCuentas.objects.create(tenant=self.tenant, cuenta='1001', descripcion='Otra caja', codigoagrupador=self.codigo)
        manual = self.manual([self.cuenta, otra])
        with patch('catalogo.views.generar_contenido_manual', return_value='Nuevo contenido') as generar:
            response = self.generar([self.cuenta.pk])
        self.assertRedirects(response, reverse('detalle_manual', args=[manual.pk]))
        generar.assert_not_called()
        self.assertEqual(self.tenant.manuales.count(), 1)

    def test_no_compara_solo_el_titulo(self):
        self.manual([], titulo='Manual: 1000')
        with patch('catalogo.views.generar_contenido_manual', return_value='Nuevo contenido'):
            self.generar([self.cuenta.pk])
        self.assertEqual(self.tenant.manuales.count(), 2)

    def test_no_duplica_manuales_sat_sin_catalogo_propio(self):
        self.cuenta.delete()
        manual = self.manual([])
        manual.codigos_sat.add(self.codigo)
        with patch('catalogo.views.generar_contenido_manual') as generar:
            response = self.generar([f'sat:{self.codigo.pk}'])
        self.assertRedirects(response, reverse('detalle_manual', args=[manual.pk]))
        generar.assert_not_called()

    def test_los_manuales_de_otra_organizacion_no_bloquean(self):
        otra = Tenant.objects.create(nombre='Otra organización', propietario=self.user)
        self.manual([], tenant=otra).codigos_sat.add(self.codigo)
        self.cuenta.delete()
        with patch('catalogo.views.generar_contenido_manual', return_value='Nuevo contenido'):
            self.generar([f'sat:{self.codigo.pk}'])
        self.assertEqual(self.tenant.manuales.count(), 1)
        self.assertEqual(otra.manuales.count(), 1)

    def test_revalida_si_otro_envio_guarda_durante_la_generacion(self):
        saved = []
        def generar(*args):
            saved.append(self.manual())
            return 'Otro contenido'
        with patch('catalogo.views.generar_contenido_manual', side_effect=generar):
            response = self.generar([self.cuenta.pk])
        self.assertRedirects(response, reverse('detalle_manual', args=[saved[0].pk]))
        self.assertEqual(self.tenant.manuales.count(), 1)

    def test_duplicados_previos_se_conservan_y_se_prefiere_el_aprobado(self):
        aprobado = self.manual(aprobado=True)
        self.manual()
        self.assertEqual(buscar_manual_existente(self.tenant, [self.cuenta], False), aprobado)
        aprobado.delete()
        self.assertIsNotNone(buscar_manual_existente(self.tenant, [self.cuenta], False))

    def test_aprobado_en_manual_conjunto_bloquea_y_no_reserva_ia(self):
        from django.contrib.messages import get_messages, WARNING
        otra = CatalogoCuentas.objects.create(tenant=self.tenant, cuenta='1001', descripcion='Otra', codigoagrupador=self.codigo)
        aprobado = self.manual([self.cuenta, otra], aprobado=True)
        self.manual()
        with patch('catalogo.views.generar_contenido_manual') as generar, patch('panel.controls.reservar_generacion') as reservar:
            response = self.generar([self.cuenta.pk])
        self.assertRedirects(response, reverse('detalle_manual', args=[aprobado.pk]))
        generar.assert_not_called(); reservar.assert_not_called()
        avisos = list(get_messages(response.wsgi_request))
        self.assertTrue(any(m.level == WARNING and 'aprobado' in str(m) for m in avisos))
        estado = self.client.get(reverse('generar_manual')).context['estados_manuales'][str(self.cuenta.pk)]
        self.assertEqual(estado, {'estado': 'aprobado', 'url': reverse('detalle_manual', args=[aprobado.pk])})

    def test_borrador_advierte_y_abre_existente(self):
        from django.contrib.messages import get_messages, WARNING
        borrador = self.manual()
        with patch('catalogo.views.generar_contenido_manual') as generar:
            response = self.generar([self.cuenta.pk])
        self.assertRedirects(response, reverse('detalle_manual', args=[borrador.pk]))
        generar.assert_not_called()
        self.assertTrue(any(m.level == WARNING and 'borrador' in str(m) for m in get_messages(response.wsgi_request)))
        self.assertEqual(self.client.get(reverse('generar_manual')).context['estados_manuales'][str(self.cuenta.pk)]['estado'], 'borrador')

    def test_aprobado_sat_se_reconoce_al_usar_catalogo_propio(self):
        aprobado = self.manual([], aprobado=True)
        aprobado.codigos_sat.add(self.codigo)
        with patch('catalogo.views.generar_contenido_manual') as generar:
            response = self.generar([self.cuenta.pk])
        self.assertRedirects(response, reverse('detalle_manual', args=[aprobado.pk]))
        generar.assert_not_called()
        self.assertEqual(self.client.get(reverse('generar_manual')).context['estados_manuales'][str(self.cuenta.pk)]['estado'], 'aprobado')

    def test_cuentas_propias_distintas_con_mismo_sat_no_se_confunden(self):
        otra = CatalogoCuentas.objects.create(tenant=self.tenant, cuenta='1001', descripcion='Otra', codigoagrupador=self.codigo)
        self.manual([otra], aprobado=True)
        with patch('catalogo.views.generar_contenido_manual', return_value='Nuevo contenido') as generar:
            self.generar([self.cuenta.pk])
        generar.assert_called_once()
        self.assertEqual(self.tenant.manuales.count(), 2)

    def test_indice_no_revela_manuales_de_otra_organizacion(self):
        otra = Tenant.objects.create(nombre='Privada', propietario=self.user)
        self.manual([self.cuenta], tenant=otra, aprobado=True)
        self.assertEqual(self.client.get(reverse('generar_manual')).context['estados_manuales'], {})

    def test_dashboard_muestra_alcance_autor_y_acciones(self):
        manual = self.manual()
        response = self.client.get(reverse('dashboard'))
        for text in ['Caja general', '1 cuenta del catálogo propio', 'Control del efectivo.', self.user.username,
                     'Editar', 'Exportar Word', 'Imprimir', 'Eliminar', 'origen']:
            self.assertContains(response, text)
        for route in ['editar_manual', 'exportar_manual', 'eliminar_manual']:
            self.assertContains(response, reverse(route, args=[manual.pk]))
        self.assertContains(response, '?imprimir=1')

    def test_editar_guarda_y_no_permite_duplicar_aunque_cambie_el_titulo(self):
        manual = self.manual()
        response = self.client.post(reverse('editar_manual', args=[manual.pk]), {
            'titulo': 'Políticas de tesorería', 'contenido': '# Control de caja actualizado'})
        self.assertRedirects(response, reverse('detalle_manual', args=[manual.pk]))
        manual.refresh_from_db()
        self.assertEqual(manual.titulo, 'Políticas de tesorería')
        self.assertEqual(manual.contenido, '# Control de caja actualizado')
        with patch('catalogo.views.generar_contenido_manual') as generar:
            response = self.generar([self.cuenta.pk])
        self.assertRedirects(response, reverse('detalle_manual', args=[manual.pk]))
        generar.assert_not_called()

    def test_exportacion_incluye_organizacion_y_sanitiza_contenido(self):
        manual = self.manual(contenido='# Título\n\n<script>alert(1)</script>\n\n**Control de caja**')
        response = self.client.get(reverse('exportar_manual', args=[manual.pk]))
        self.assertEqual(response.status_code, 200)
        from io import BytesIO
        from zipfile import ZipFile
        from docx import Document
        self.assertTrue(response['Content-Disposition'].endswith('.docx"'))
        self.assertIn('wordprocessingml.document', response['Content-Type'])
        documento = Document(BytesIO(response.content))
        texto = '\n'.join(p.text for p in documento.paragraphs)
        cabecera = '\n'.join(c.text for t in documento.sections[0].header.tables for f in t.rows for c in f.cells)
        self.assertIn(self.tenant.nombre, cabecera)
        self.assertIn('Control de caja', texto)
        with ZipFile(BytesIO(response.content)) as archivo:
            xml = archivo.read('word/document.xml').decode()
        self.assertIn('<w:b', xml)
        self.assertNotIn('<script>', xml)

    def test_impresion_automatica_solo_con_parametro_explicito(self):
        manual = self.manual()
        url = reverse('detalle_manual', args=[manual.pk])
        self.assertNotContains(self.client.get(url), 'await document.fonts.ready')
        self.assertContains(self.client.get(url + '?imprimir=1'), 'await document.fonts.ready')

    def test_eliminar_es_post_y_regresa_al_dashboard(self):
        manual = self.manual()
        url = reverse('eliminar_manual', args=[manual.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        response = self.client.post(url, {'origen': 'dashboard'})
        self.assertRedirects(response, reverse('dashboard'))
        self.assertFalse(ManualContable.objects.filter(pk=manual.pk).exists())

    def test_eliminar_requiere_csrf(self):
        manual = self.manual()
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        session = client.session
        session['tenant_id'] = self.tenant.pk
        session.save()
        self.assertEqual(client.post(reverse('eliminar_manual', args=[manual.pk])).status_code, 403)
        self.assertTrue(ManualContable.objects.filter(pk=manual.pk).exists())

    def test_acciones_no_acceden_a_manuales_de_otra_organizacion(self):
        otra = Tenant.objects.create(nombre='Privada', propietario=self.user)
        manual = self.manual([], tenant=otra)
        for route in ['detalle_manual', 'editar_manual', 'exportar_manual']:
            self.assertEqual(self.client.get(reverse(route, args=[manual.pk])).status_code, 404)
        for route in ['eliminar_manual', 'editar_manual']:
            self.assertEqual(self.client.post(reverse(route, args=[manual.pk]), {'titulo': 'Intento'}).status_code, 404)
        self.assertNotContains(self.client.get(reverse('dashboard')), manual.titulo)
